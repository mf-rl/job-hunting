#!/usr/bin/env python3
"""Render tailored CV JSON into the uploaded CV template style."""
from __future__ import annotations

import argparse
import json
import os
import subprocess
import sys
from pathlib import Path

try:
    from docx import Document
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    print("python-docx is required. Install with: pip install python-docx", file=sys.stderr)
    sys.exit(1)

FONT_NAME = "Arial"
TEXT_COLOR = RGBColor(0x11, 0x11, 0x11)
MUTED_COLOR = RGBColor(0x28, 0x28, 0x28)
SECTION_COLOR = RGBColor(0x00, 0x00, 0x00)

MONTHS = {
    "01": "Jan",
    "02": "Feb",
    "03": "Mar",
    "04": "Apr",
    "05": "May",
    "06": "Jun",
    "07": "Jul",
    "08": "Aug",
    "09": "Sep",
    "10": "Oct",
    "11": "Nov",
    "12": "Dec",
}


def _text(value: object) -> str:
    return str(value or "").replace("\u00a0", " ").strip()


def _items(values: object) -> list[str]:
    if not isinstance(values, list):
        return []
    return [_text(value) for value in values if _text(value)]


def _format_date(value: object) -> str:
    raw = _text(value)
    if not raw:
        return ""
    lowered = raw.lower()
    if lowered in {"present", "current", "now"}:
        return "Present"
    if len(raw) >= 7 and raw[4] == "-":
        year, month = raw[:4], raw[5:7]
        if year.isdigit() and month in MONTHS:
            return f"{MONTHS[month]} {year}"
    return raw


def _date_range(start: object, end: object) -> str:
    parts = [_format_date(start), _format_date(end)]
    parts = [part for part in parts if part]
    return " - ".join(parts)


def _font(run, size: float, *, bold: bool = False, italic: bool = False, color: RGBColor = TEXT_COLOR) -> None:
    run.bold = bold
    run.italic = italic
    run.font.name = FONT_NAME
    run.font.size = Pt(size)
    run.font.color.rgb = color
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _paragraph_format(paragraph, *, before: float = 0, after: float = 0, line_spacing: float = 1.0) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line_spacing


def _add_text(doc: Document, text: str, *, size: float = 8.8, bold: bool = False,
              italic: bool = False, before: float = 0, after: float = 0,
              color: RGBColor = TEXT_COLOR, align: int | None = None):
    paragraph = doc.add_paragraph()
    if align is not None:
        paragraph.alignment = align
    _paragraph_format(paragraph, before=before, after=after)
    run = paragraph.add_run(text)
    _font(run, size, bold=bold, italic=italic, color=color)
    return paragraph


def _add_section(doc: Document, title: str) -> None:
    _add_text(doc, title.upper(), size=9.4, bold=True, before=8, after=3, color=SECTION_COLOR)


def _add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    _paragraph_format(paragraph, after=1, line_spacing=1.0)
    paragraph.paragraph_format.left_indent = Inches(0.25)
    paragraph.paragraph_format.first_line_indent = Inches(-0.14)
    marker = paragraph.add_run("\u2022 ")
    _font(marker, 8.4, color=TEXT_COLOR)
    run = paragraph.add_run(text)
    _font(run, 8.4, color=TEXT_COLOR)


def _configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.45)
    section.bottom_margin = Inches(0.45)
    section.left_margin = Inches(0.55)
    section.right_margin = Inches(0.55)

    styles = doc.styles
    for style_name in ("Normal", "Body Text", "List Bullet"):
        try:
            style = styles[style_name]
        except KeyError:
            continue
        style.font.name = FONT_NAME
        style.font.size = Pt(8.8)
        style.font.color.rgb = TEXT_COLOR
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_NAME)


def _contact_line(contact: object) -> str:
    if not isinstance(contact, dict):
        return ""
    values = []
    for key in ("phone", "email", "linkedin", "location"):
        value = _text(contact.get(key))
        if value:
            values.append(value)
    return " - ".join(values)


def _split_experience(experiences: object) -> tuple[list[dict], list[dict]]:
    detailed: list[dict] = []
    earlier: list[dict] = []
    if not isinstance(experiences, list):
        return detailed, earlier
    for item in experiences:
        if not isinstance(item, dict):
            continue
        bullets = _items(item.get("bullets"))
        if bullets:
            detailed.append(item)
        else:
            earlier.append(item)
    return detailed, earlier


def _experience_header(exp: dict) -> tuple[str, str]:
    company = _text(exp.get("company"))
    location = _text(exp.get("location"))
    role = _text(exp.get("role"))
    dates = _date_range(exp.get("start"), exp.get("end"))

    company_line = company
    if company and location:
        company_line = f"{company} - {location}"
    elif location:
        company_line = location

    role_line = role
    if role and dates:
        role_line = f"{role} | {dates}"
    elif dates:
        role_line = dates
    return company_line, role_line


def _render_experience(doc: Document, experiences: list[dict]) -> None:
    _add_section(doc, "Professional Experience")
    for index, exp in enumerate(experiences):
        company_line, role_line = _experience_header(exp)
        if company_line:
            _add_text(doc, company_line, size=8.8, bold=True, before=3 if index else 0, after=0)
        if role_line:
            _add_text(doc, role_line, size=8.4, after=2, color=MUTED_COLOR)
        for bullet in _items(exp.get("bullets")):
            _add_bullet(doc, bullet)


def _render_earlier_experience(doc: Document, experiences: list[dict]) -> None:
    if not experiences:
        return
    _add_section(doc, "Earlier Experience")
    _add_text(
        doc,
        "Delivered multiple enterprise and financial system solutions across the following roles:",
        size=8.6,
        after=2,
    )
    for exp in experiences:
        company = _text(exp.get("company"))
        role = _text(exp.get("role"))
        dates = _date_range(exp.get("start"), exp.get("end"))
        line = company
        if company and role:
            line = f"{company} - {role}"
        elif role:
            line = role
        if dates:
            line = f"{line} ({dates})" if line else dates
        if line:
            _add_text(doc, line, size=8.5, after=0)


def _render_education(doc: Document, education: object) -> None:
    if not isinstance(education, list) or not education:
        return
    _add_section(doc, "Education")
    for edu in education:
        if not isinstance(edu, dict):
            continue
        degree = _text(edu.get("degree"))
        field = _text(edu.get("field"))
        institution = _text(edu.get("institution"))
        year = _text(edu.get("year") or edu.get("end"))

        credential = degree
        if degree and field:
            credential = f"{degree} in {field}"
        elif field:
            credential = field

        line = credential or institution
        if institution and institution != line:
            line = f"{line} - {institution}" if line else institution
        if year:
            line = f"{line} ({year})" if line else year
        if line:
            _add_text(doc, line, size=8.6, after=1)


def _render_skills(doc: Document, skills: object) -> None:
    if not isinstance(skills, list) or not skills:
        return
    _add_section(doc, "Technical Skills")
    for group in skills:
        if not isinstance(group, dict):
            continue
        category = _text(group.get("category"))
        items = _items(group.get("items"))
        if not category and not items:
            continue
        paragraph = doc.add_paragraph()
        _paragraph_format(paragraph, after=1)
        if category:
            label = paragraph.add_run(f"{category}: ")
            _font(label, 8.5, bold=True)
        values = paragraph.add_run(", ".join(items))
        _font(values, 8.5)


def _render_additional_information(doc: Document, cv: dict) -> None:
    lines: list[tuple[str, str]] = []

    languages = cv.get("languages")
    if isinstance(languages, list):
        parts = []
        for language in languages:
            if not isinstance(language, dict):
                continue
            name = _text(language.get("lang"))
            level = _text(language.get("level"))
            if name and level:
                parts.append(f"{name} ({level})")
            elif name:
                parts.append(name)
        if parts:
            lines.append(("Languages", ", ".join(parts)))

    competencies = _items(cv.get("competencies"))
    if competencies:
        lines.append(("Core Competencies", ", ".join(competencies)))

    additional = cv.get("additional_information") or cv.get("additionalInfo")
    if isinstance(additional, dict):
        for key, value in additional.items():
            label = str(key).replace("_", " ").title()
            if isinstance(value, list):
                rendered_value = ", ".join(_items(value))
            else:
                rendered_value = _text(value)
            if rendered_value:
                lines.append((label, rendered_value))

    if not lines:
        return
    _add_section(doc, "Additional Information")
    for label, value in lines:
        paragraph = doc.add_paragraph()
        _paragraph_format(paragraph, after=1)
        label_run = paragraph.add_run(f"{label}: ")
        _font(label_run, 8.5, bold=True)
        value_run = paragraph.add_run(value)
        _font(value_run, 8.5)


def render_docx(cv: dict, output_path: str) -> str:
    """Render a tailored CV JSON to a DOCX file using the template layout."""
    doc = Document()
    _configure_document(doc)

    name = _text(cv.get("name")) or "Your Name"
    headline = _text(cv.get("headline"))
    contact = _contact_line(cv.get("contact"))

    _add_text(doc, name, size=13.5, bold=True, after=1)
    if headline:
        _add_text(doc, headline, size=8.8, after=1)
    if contact:
        _add_text(doc, contact, size=8.4, after=5)

    summary = _text(cv.get("summary"))
    if summary:
        _add_section(doc, "Professional Summary")
        _add_text(doc, summary, size=8.8, after=2)

    detailed, earlier = _split_experience(cv.get("experience"))
    if detailed:
        _render_experience(doc, detailed)
    if earlier:
        _render_earlier_experience(doc, earlier)

    _render_education(doc, cv.get("education"))
    _render_skills(doc, cv.get("skills"))
    _render_additional_information(doc, cv)

    doc.save(output_path)
    return output_path


def convert_to_pdf(docx_path: str, output_dir: str | None = None) -> str | None:
    """Use LibreOffice headless to convert a DOCX file to PDF when available."""
    if output_dir is None:
        output_dir = os.path.dirname(docx_path)

    try:
        result = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", output_dir, docx_path],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if result.returncode == 0:
            pdf_path = str(Path(docx_path).with_suffix(".pdf"))
            if os.path.exists(pdf_path):
                return pdf_path
            expected = Path(output_dir) / Path(docx_path).with_suffix(".pdf").name
            if expected.exists():
                return str(expected)
        print(f"LibreOffice warning: {result.stderr[:200]}", file=sys.stderr)
    except FileNotFoundError:
        print("soffice not found - install LibreOffice for DOCX to PDF conversion", file=sys.stderr)
    except Exception as exc:  # pragma: no cover - defensive CLI boundary
        print(f"PDF conversion failed: {exc}", file=sys.stderr)
    return None


def main() -> None:
    parser = argparse.ArgumentParser(description="Render tailored CV JSON to DOCX")
    parser.add_argument("cv_json", help="Path to tailored CV JSON file")
    parser.add_argument("--output", "-o", help="Output path for DOCX (default: <input>.docx)")
    parser.add_argument("--pdf", action="store_true", help="Also convert to PDF with LibreOffice when available")
    args = parser.parse_args()

    with open(args.cv_json, encoding="utf-8") as fh:
        cv = json.load(fh)

    output = args.output or str(Path(args.cv_json).with_suffix(".docx"))
    render_docx(cv, output)
    print(json.dumps({"status": "ok", "docx": output, "size": os.path.getsize(output)}))

    if args.pdf:
        pdf = convert_to_pdf(output)
        if pdf:
            print(json.dumps({"status": "ok", "pdf": pdf, "size": os.path.getsize(pdf)}))


if __name__ == "__main__":
    main()
